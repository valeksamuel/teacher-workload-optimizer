from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

import constants as c
from enums import Week
from i18n import translate
from models import Assignment, Teacher, TeacherStat


def export_to_excel(
    output_path: str,
    assignments1: list[Assignment],
    stats1: list[TeacherStat],
    assignments2: list[Assignment],
    stats2: list[TeacherStat],
    teachers_by_id: dict[str, Teacher],
) -> None:
    wb = Workbook()
    default_sheet = wb.active
    assert default_sheet is not None
    wb.remove(default_sheet)

    sheet_assignments(
        wb, translate("excel.sheet1.header"), assignments1, teachers_by_id, 1
    )
    sheet_teachers(wb, translate("excel.sheet2.header"), stats1, 1)
    sheet_assignments(
        wb, translate("excel.sheet3.header"), assignments2, teachers_by_id, 2
    )
    sheet_teachers(wb, translate("excel.sheet4.header"), stats2, 2)

    wb.save(output_path)

    print(translate("excel.export.complete", output_path=output_path))


def sheet_assignments(
    wb: Workbook,
    title: str,
    assignments: list[Assignment],
    teachers_by_id: dict[str, Teacher],
    phase_num: int,
) -> None:
    ws = wb.create_sheet(
        title=translate("excel.sheet.assignments.title", phase_num=phase_num)
    )
    title_row(ws, title, len(c.LABELS_SHEET_ASSIGNMENTS))
    header_row(ws, 2, c.LABELS_SHEET_ASSIGNMENTS)

    for i, a in enumerate(assignments):
        r = i + 3
        ws.cell(r, 1, a.lesson.id)
        ws.cell(r, 2, a.lesson.class_name)
        ws.cell(r, 3, a.lesson.get_subject_abbreviation)
        ws.cell(r, 4, a.lesson.two_week_hours)
        ws.cell(r, 5, a.lesson.week)
        assert a.assigned_teacher_id is not None, (
            "Assignment must have an assigned teacher"
        )
        ws.cell(r, 6, teachers_by_id[a.assigned_teacher_id].name)
        ws.cell(r, 4).alignment = Alignment(horizontal="center")

        if a.assigned_teacher_id is None:
            fill = c.ERR_FILL
        elif a.assigned_teacher_id == Week.EVEN:
            fill = c.BLUE_FILL
        elif a.assigned_teacher_id == Week.ODD:
            fill = c.GOLD_FILL
        else:
            fill = c.GRN_FILL

        for j in range(1, 7):
            ws.cell(r, j).fill = fill
            ws.cell(r, j).font = c.NORM

    auto_width(ws)
    ws.freeze_panes = "A3"

    r_leg = len(assignments) + 4
    ws.cell(r_leg, 1, translate("excel.legend.name")).font = c.BOLD
    ws.cell(r_leg + 1, 1, translate("excel.legend.even")).fill = c.BLUE_FILL
    ws.cell(r_leg + 2, 1, translate("excel.legend.odd")).fill = c.GOLD_FILL
    ws.cell(r_leg + 3, 1, translate("excel.legend.every")).fill = c.GRN_FILL
    ws.cell(r_leg + 4, 1, translate("excel.legend.not.assigned")).fill = c.ERR_FILL

    for dr in range(5):
        ws.cell(r_leg + dr, 1).font = c.NORM


def title_row(ws: Worksheet, text: str, num_cols: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
    cell = ws.cell(row=1, column=1, value=text)
    cell.fill = c.TITLE_ROW
    cell.font = c.TITLE_FNT
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24


def header_row(ws: Worksheet, row: int, labels: list[str]) -> None:
    for j, lbl in enumerate(labels, 1):
        cell = ws.cell(row=row, column=j, value=lbl)
        cell.fill = c.HDR_FILL
        cell.font = c.HDR_FONT
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    ws.row_dimensions[row].height = 28


def auto_width(ws: Worksheet) -> None:
    for col_idx, col in enumerate(ws.columns, start=1):
        max_len = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 42)


def sheet_teachers(
    wb: Workbook, title: str, stats: list[TeacherStat], phase_num: int
) -> None:
    ws = wb.create_sheet(
        title=translate("excel.sheet.teachers.title", phase_num=phase_num)
    )
    title_row(ws, title, len(c.LABELS_SHEET_TEACHERS))
    header_row(ws, 2, c.LABELS_SHEET_TEACHERS)

    for i, s in enumerate(stats):
        r = i + 3
        vals = [
            s.teacher,
            s.teacher.min_two_week_workload,
            s.teacher.target_two_week_workload,
            s.teacher.max_two_week_workload,
            s.total_workload,
            s.even_workload,
            s.odd_workload,
            s.target_deviation,
            s.even_odd_diff,
        ]

        for j, v in enumerate(vals, 1):
            cell = ws.cell(r, j, v)
            cell.font = c.NORM
            cell.alignment = Alignment(horizontal="center" if j > 1 else "left")

        total = s.total_workload
        min_w = s.teacher.min_two_week_workload
        max_w = s.teacher.max_two_week_workload
        dev = abs(s.target_deviation)

        if not (min_w <= total <= max_w):  # not possible
            row_fill = c.ERR_FILL
        elif dev <= 2:  # replace for constants
            row_fill = c.OK_FILL
        elif dev <= 5:  # replace for constants
            row_fill = c.WARN_FILL
        else:
            row_fill = c.ALT_FILL

        for j in range(1, len(c.LABELS_SHEET_TEACHERS) + 1):
            ws.cell(r, j).fill = row_fill

        # Highlighting of the deviation cell
        c_dev = ws.cell(r, 8)
        c_dev.fill = (
            c.OK_FILL if dev <= 2 else c.WARN_FILL if dev <= 5 else c.ERR_FILL
        )  # replace for constants

        # Highlighting of the E-O difference cell
        rd = abs(s.even_odd_diff)
        c_diff = ws.cell(r, 9)
        c_diff.fill = (
            c.OK_FILL if rd <= 1 else c.WARN_FILL if rd <= 4 else c.ERR_FILL
        )  # replace for constants

    # Summary row
    r_sum = len(stats) + 3
    ws.cell(r_sum, 1, translate("summary.sum")).font = c.BOLD
    ws.cell(r_sum, 5, translate("summary.sum.value", r_sum=r_sum)).font = c.BOLD
    total_dev = sum(abs(s.target_deviation) for s in stats)
    total_diff = sum(abs(s.even_odd_diff) for s in stats)
    ws.cell(r_sum, 8, round(total_dev, 1)).font = c.BOLD
    ws.cell(r_sum, 9, round(total_diff, 1)).font = c.BOLD
    for j in range(1, len(c.LABELS_SHEET_TEACHERS) + 1):
        ws.cell(r_sum, j).fill = c.HDR_FILL
        if not ws.cell(r_sum, j).font.bold:
            ws.cell(r_sum, j).font = c.HDR_FONT

    # Color legend
    r_leg = r_sum + 2
    ws.cell(r_leg, 1, translate("excel.legend.name")).font = c.BOLD
    legend_items = [
        (c.OK_FILL, translate("excel.legend.text.ok")),  # replace for constants
        (c.WARN_FILL, translate("excel.legend.text.warn")),  # replace for constants
        (c.ALT_FILL, translate("excel.legend.text.alt")),  # replace for constants
        (c.ERR_FILL, translate("excel.legend.text.err")),
    ]

    for k, (fill, text) in enumerate(legend_items):
        ws.cell(r_leg + 1 + k, 1, text).fill = fill
        ws.cell(r_leg + 1 + k, 1).font = c.NORM

    auto_width(ws)
    ws.freeze_panes = "A3"
