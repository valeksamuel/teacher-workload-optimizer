from openpyxl.styles import Font, PatternFill

from i18n import translate

COL_LESSON_ID = translate("excel.col.lesson.id")
COL_CLASS_NAME = translate("excel.col.class.name")
COL_LESSON_NAME = translate("excel.col.lesson.name")
COL_LESSON_WEEKLY_LOAD = translate("excel.col.lesson.weekly.load")
COL_LESSON_WEEK = translate("excel.col.lesson.week")
COL_SUBJECT_TEACHER = translate("excel.col.subject.teacher")
COL_TEACHER_ID = translate("excel.col.teacher.id")
COL_TEACHER_NAME = translate("excel.col.teacher.name")
COL_TEACHER_MIN_2WEEKS_WORKLOAD = translate("excel.col.min.2week.workload")
COL_TEACHER_TARGET_2WEEKS_WORKLOAD = translate("excel.col.target.2week.workload")
COL_TEACHER_MAX_2WEEKS_WORKLOAD = translate("excel.col.max.2week.workload")
INFO_COLS = {
    COL_TEACHER_ID,
    COL_TEACHER_NAME,
    COL_TEACHER_MIN_2WEEKS_WORKLOAD,
    COL_TEACHER_TARGET_2WEEKS_WORKLOAD,
    COL_TEACHER_MAX_2WEEKS_WORKLOAD,
}
WEEK_TYPE_EVEN = translate("week.type.even")
WEEK_TYPE_ODD = translate("week.type.odd")
WEEK_TYPE_EVERY = translate("week.type.every")

# Console printing - adjust column widths and alignment for better readability
SEPARATOR = "─" * 62
CONSOLE_COLS = [
    (translate("console.col.teacher"), 30, "<"),
    (translate("console.col.workload"), 7, ">"),
    (translate("console.col.target"), 5, ">"),
    (translate("console.col.deviation"), 6, ">"),
    (translate("console.col.even"), 7, ">"),
    (translate("console.col.odd"), 7, ">"),
    (translate("console.col.e.o.diff"), 6, ">"),
]

# Excel output
LABELS_SHEET_ASSIGNMENTS = [
    translate("excel.label.id"),
    translate("excel.label.class"),
    translate("excel.label.subject"),
    translate("excel.label.2week.workload"),
    translate("excel.label.week"),
    translate("excel.label.teacher"),
]
LABELS_SHEET_TEACHERS = [
    translate("excel.label.min"),
    translate("excel.label.target"),
    translate("excel.label.max"),
    translate("excel.label.total.workload"),
    translate("excel.label.even.week"),
    translate("excel.label.odd.week"),
    translate("excel.label.deviation"),
    translate("excel.label.difference"),
]
HDR_FILL = PatternFill("solid", fgColor="1F3864")  # dark blue - header
HDR_FONT = Font(bold=True, color="FFFFFF", name="Arial")
ALT_FILL = PatternFill("solid", fgColor="EEF2F7")  # light blue - even rows
OK_FILL = PatternFill("solid", fgColor="D6EAD6")  # green - in limit
WARN_FILL = PatternFill("solid", fgColor="FFF3CD")  # yellow - slight deviation
ERR_FILL = PatternFill("solid", fgColor="FADADD")  # red - out of limit
BLUE_FILL = PatternFill("solid", fgColor="D1E8F5")  # blue - even week
GOLD_FILL = PatternFill("solid", fgColor="FFF0B3")  # gold - odd week
GRN_FILL = PatternFill("solid", fgColor="E0F5E0")  # light green - every week
TITLE_ROW = PatternFill("solid", fgColor="2E4057")
TITLE_FNT = Font(bold=True, color="FFFFFF", size=12, name="Arial")
BOLD = Font(bold=True, name="Arial")
NORM = Font(name="Arial")
